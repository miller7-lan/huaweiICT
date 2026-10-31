"""Build the development design document draft as .docx.

Chapters 1-3 are written; chapters 4-8 are a framework of headings, prefilled
facts and yellow placeholders. Needs python-docx, which is not in the project
lock; run it from a separate environment. Every number comes from the dated
evidence named in docs/2026-10-18对话决策与接续记录.md sections 25-46.
Once teammates edit the .docx, do not regenerate over it without merging.
Output: docs/competition/开发设计文档-20261031.docx
"""
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_COLOR_INDEX
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
FIGURES = ROOT / "docs/competition/assets/design"
SCREENSHOTS = ROOT / "docs/competition/assets/screenshots"
OUTPUT = ROOT / "docs/competition/开发设计文档-20261031.docx"
BODY_FONT, HEAD_FONT, LATIN = "宋体", "黑体", "Times New Roman"


def set_fonts(target, east, size=None, bold=None, color=True):
    """Apply Latin and East Asian fonts; drop theme fonts that would override them."""
    font = target.font
    font.name = LATIN
    if size:
        font.size = Pt(size)
    if bold is not None:
        font.bold = bold
    if color:
        font.color.rgb = RGBColor(0, 0, 0)
    rpr = target.element.get_or_add_rPr()
    fonts = rpr.find(qn("w:rFonts"))
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        rpr.insert(0, fonts)
    for attribute in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
        fonts.attrib.pop(qn(attribute), None)
    for attribute, value in (("w:ascii", LATIN), ("w:hAnsi", LATIN), ("w:eastAsia", east)):
        fonts.set(qn(attribute), value)


def configure(doc):
    section = doc.sections[0]
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    section.top_margin = section.bottom_margin = Cm(2.54)
    section.left_margin = section.right_margin = Cm(3.17)
    normal = doc.styles["Normal"]
    set_fonts(normal, BODY_FONT, 12)
    normal.paragraph_format.line_spacing = 1.5
    normal.paragraph_format.first_line_indent = Pt(24)
    normal.paragraph_format.space_after = Pt(0)
    for name, size, before, after in (("Heading 1", 16, 18, 12), ("Heading 2", 14, 12, 6), ("Heading 3", 12, 6, 3)):
        style = doc.styles[name]
        set_fonts(style, HEAD_FONT, size, True)
        style.paragraph_format.first_line_indent = Pt(0)
        style.paragraph_format.space_before, style.paragraph_format.space_after = Pt(before), Pt(after)
        style.paragraph_format.line_spacing = 1.3
        style.paragraph_format.keep_with_next = True
    bullet = doc.styles["List Bullet"]
    set_fonts(bullet, BODY_FONT, 12)
    bullet.paragraph_format.first_line_indent = Pt(0)
    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.CENTER
    header.paragraph_format.first_line_indent = Pt(0)
    set_fonts(header.add_run("政务资产认知与可进化决策智能体 · 开发设计文档"), BODY_FONT, 9)
    border = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    for key, value in (("w:val", "single"), ("w:sz", "4"), ("w:space", "1"), ("w:color", "808080")):
        bottom.set(qn(key), value)
    border.append(bottom)
    header._p.get_or_add_pPr().append(border)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.paragraph_format.first_line_indent = Pt(0)
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    run = OxmlElement("w:r")
    text = OxmlElement("w:t")
    text.text = "1"
    run.append(text)
    field.append(run)
    footer._p.append(field)


class Writer:
    def __init__(self, doc):
        self.doc = doc

    def heading(self, text, level):
        self.doc.add_heading(text, level=level)

    def para(self, *parts, indent=True, align=None, size=None, east=BODY_FONT):
        """parts: plain strings, or (text, bold) pairs."""
        paragraph = self.doc.add_paragraph()
        if not indent:
            paragraph.paragraph_format.first_line_indent = Pt(0)
        if align:
            paragraph.alignment = align
        for part in parts:
            text, bold = (part, False) if isinstance(part, str) else part
            set_fonts(paragraph.add_run(text), east, size, bold)
        return paragraph

    def lead(self, label, text):
        return self.para((label, True), text)

    def bullets(self, items):
        for item in items:
            paragraph = self.doc.add_paragraph(style="List Bullet")
            if isinstance(item, tuple):
                set_fonts(paragraph.add_run(item[0]), BODY_FONT, None, True)
                set_fonts(paragraph.add_run(item[1]), BODY_FONT)
            else:
                set_fonts(paragraph.add_run(item), BODY_FONT)

    def caption(self, text):
        paragraph = self.para(text, indent=False, align=WD_ALIGN_PARAGRAPH.CENTER, size=10.5, east=HEAD_FONT)
        paragraph.paragraph_format.space_before = Pt(3)
        paragraph.paragraph_format.space_after = Pt(6)
        return paragraph

    def table(self, caption, headers, rows, widths):
        self.caption(caption).paragraph_format.keep_with_next = True
        table = self.doc.add_table(rows=1 + len(rows), cols=len(headers))
        table.style = "Table Grid"
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = False
        # WPS and LibreOffice size columns from the grid, not from cell widths.
        for column, width in zip(table.columns, widths):
            column.width = Cm(width)
        layout = OxmlElement("w:tblLayout")
        layout.set(qn("w:type"), "fixed")
        table._tbl.tblPr.append(layout)
        for r, values in enumerate([headers] + rows):
            row = table.rows[r]
            properties = row._tr.get_or_add_trPr()
            # Rows never split across pages; the header repeats and stays with the first row.
            properties.append(OxmlElement("w:cantSplit"))
            if r == 0:
                repeat = OxmlElement("w:tblHeader")
                repeat.set(qn("w:val"), "true")
                properties.append(repeat)
            for c, value in enumerate(values):
                cell = row.cells[c]
                cell.width = Cm(widths[c])
                paragraph = cell.paragraphs[0]
                paragraph.paragraph_format.first_line_indent = Pt(0)
                paragraph.paragraph_format.line_spacing = 1.2
                paragraph.paragraph_format.keep_with_next = r == 0
                set_fonts(paragraph.add_run(value), HEAD_FONT if r == 0 else BODY_FONT, 10.5, r == 0)
                if r == 0:
                    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    shade = OxmlElement("w:shd")
                    for key, val in (("w:val", "clear"), ("w:color", "auto"), ("w:fill", "D9E2F3")):
                        shade.set(qn(key), val)
                    cell._tc.get_or_add_tcPr().append(shade)
        spacer = self.para("", indent=False)
        spacer.paragraph_format.line_spacing = 1.0

    def figure(self, filename, caption, width=15.5, folder=FIGURES):
        paragraph = self.para("", indent=False, align=WD_ALIGN_PARAGRAPH.CENTER)
        paragraph.paragraph_format.keep_with_next = True
        # Body line spacing would multiply the inline picture's height.
        paragraph.paragraph_format.line_spacing = 1.0
        paragraph.paragraph_format.space_before = Pt(6)
        paragraph.runs[0].add_picture(str(folder / filename), width=Cm(width))
        self.caption(caption)

    def todo(self, owner, points, materials=None):
        """Yellow-highlighted placeholder; deleted when the section is written."""
        text = f"【待补｜负责人：{owner}】{points}"
        if materials:
            text += f"素材：{materials}"
        paragraph = self.para(text, indent=False, size=10.5)
        paragraph.paragraph_format.line_spacing = 1.25
        paragraph.paragraph_format.space_after = Pt(4)
        for run in paragraph.runs:
            run.font.highlight_color = WD_COLOR_INDEX.YELLOW
        return paragraph


def title_block(w):
    title = w.para(("基于 Nexent 的政务资产认知与可进化决策智能体", True), indent=False,
                   align=WD_ALIGN_PARAGRAPH.CENTER, size=18, east=HEAD_FONT)
    title.paragraph_format.space_after = Pt(6)
    w.para("开发设计文档（初稿）", indent=False, align=WD_ALIGN_PARAGRAPH.CENTER, size=14, east=HEAD_FONT)
    w.para("第十一届华为 ICT 大赛中国创新赛 · 赛题 3", indent=False, align=WD_ALIGN_PARAGRAPH.CENTER, size=10.5)
    w.para("基于华为 ModelEngine Nexent 打造可进化决策智能体", indent=False, align=WD_ALIGN_PARAGRAPH.CENTER, size=10.5)
    w.para("稿件日期：2026 年 10 月 31 日。学校、队名、成员等信息按官方模板填写。", indent=False,
           align=WD_ALIGN_PARAGRAPH.CENTER, size=10.5)
    note = w.todo("全体", "黄色底纹段落为待补内容，写好对应小节后整段删除。第 1–3 章为初稿，第 4–8 章为章节框架；"
                         "已有事实预填在表格中，“待测”“待补”处需按实际结果填写。")
    note.paragraph_format.space_after = Pt(12)


def chapter_one(w):
    w.heading("1 项目概述", 1)
    w.heading("1.1 背景与痛点", 2)
    w.para("政务服务领域积累了大量公开资料：办事指南、政策问答、政策解读、收费目录表格和扫描版政策附件。"
           "以居民身份证申领、换领、补领为例，同一事项的依据分散在不同部门、不同年份、不同格式的文件中，"
           "长期处于“能下载、难利用”的沉睡状态。群众咨询和窗口复核都需要把这些资料拼起来，并说明每个结论出自哪里。"
           "本项目在真实公开资料上归纳出以下五类痛点。")
    w.lead("（1）多模态异构，资料难以联通。", "收费依据的文号记在 Excel 收费目录的单元格里，收费分支写在财政部通知的网页正文中，"
           "历史便民措施只存在于扫描版 PDF。只做文本检索，难以把三类资料连成一条依据链。")
    w.lead("（2）跨文档口径复杂，容易答错分支。", "首次申领居民身份证的工本费自 2018 年 4 月 1 日起停征（财税〔2018〕37 号），"
           "换领仍按 2007 年通知每证 20 元（财综〔2007〕34 号）。回答“换领收多少钱”必须同时读两份文件，并区分首次申领、"
           "换领、补领等分支。")
    w.lead("（3）知识持续变化，旧依据仍需追溯。", "河南身份证办理时限由 2024 年措施的 25 个自然日调整为 2026 年措施的 15 个工作日。"
           "两者单位不同，不能直接相减；新资料到来时，旧依据也要保留，以便核对历史办件。")
    w.lead("（4）标注稀缺，属于典型低资源场景。", "政务领域缺少成规模的标注数据，而条款中的对象、地区、期限和例外又必须完全准确，"
           "难以依靠大规模标注训练专用模型。")
    w.lead("（5）可追溯要求高。", "面向群众的答复必须给出原文出处；资料不足或现行效力未确认时，系统应当明确停止，而不是猜测。")

    w.heading("1.2 项目目标与定位", 2)
    w.para("本项目基于华为 ModelEngine Nexent 平台，构建政务领域的资产认知与可进化决策智能体。"
           "以河南、南京、北京等地居民身份证业务的公开资料为试点，目标是：")
    w.bullets([
        ("多模态资产统一识别与标识：", "HTML、Excel、扫描 PDF 统一登记原件哈希、修订版本和精确定位，并转为可检索、可引用的资产单元。"),
        ("低资源半自动本体构建：", "从少量原文语法种子生成候选类型、关系和条件，人工确认结构后成为可执行的推理计划，并随资料变化持续演化。"),
        ("检索-推理双驱动执行流：", "借助 Nexent 的主控与子 Agent 编排、MCP 工具的确定性执行，完成跨文档多跳推理，每一步都可溯源。"),
        ("可复用 Skill 工作流模板：", "把验证过的编排模式沉淀为参数化模板，新事项、新地区按规格即可快速部署。"),
    ])
    w.para("系统服务三类用户：咨询群众获得带原文出处的答复；窗口工作人员按固定知识版本复核和交接；"
           "资料维护者登记来源、审核更新候选、管理知识版本。")

    w.heading("1.3 主要创新点", 2)
    w.lead("（1）多模态资产统一标识与跨格式推理。", "每份原件以 SHA-256 标识，修订与定位分开记录：HTML 记到段落和条文，"
           "Excel 记到单元格和行，扫描 PDF 记到 PDF 页码及左右半页。河南试点中，Excel、HTML、扫描 PDF 三种格式的"
           " 17 处定位共同进入同一条四节点推理计划，表格提供规范文号，网页提供收费分支，扫描件提供历史服务条件。")
    w.lead("（2）低资源半自动本体构建与演化。", "不依赖大规模标注：少量原文语法种子从完整条款中抽取候选，人工只需确认结构；"
           "确认后的关系成为带类型、带来源、带依赖的执行计划。版本不可变、可一键回退，来源发生变化时旧结果自动阻断。"
           "北京、南京、河南三地分别完成了新类型或新关系的构建与执行。")
    w.lead("（3）检索-推理双驱动与确定性核验。", "模型只负责任务调度和答复组织；检索、本体计划执行、引文逐字核对、数值与单位核对"
           "都由 MCP 工具完成。任何环节不通过即停止作答，并说明缺少什么依据。")
    w.lead("（4）可复用 Skill 工作流模板与自主更新。", "把“主控—检索—核验”的编排抽象为分层模板，由 profile 规格生成 Skill 和主控协议，"
           "河南链路的 Skill 可由模板逐字复现。主控心跳与官方来源观察发现新资料后，交由整理和核验子 Agent 生成待审更新候选。")

    w.heading("1.4 当前完成情况", 2)
    w.para("截至 2026 年 10 月 30 日，主要能力的实现与验证结果见表 1-1。其中南京、河南两条推理链、"
           "河南自主更新和自由提问问答均在本机 Nexent v2.7.0 上实际运行。")
    w.table("表 1-1 主要能力与验证结果", ["能力", "实现与验证结果"], [
        ["多模态冻结语料", "7 份官方原件（HTML、Excel、扫描 PDF）整理为 581 个检索单元，原件 SHA-256 全部核对一致；"
                     "资料目录累计登记 79 份文档、186 个修订（10 月 18 日快照）"],
        ["跨文档多跳推理（南京）", "本体控制的三步依赖计划在 Nexent 主控与两个子 Agent 中实跑，13 项验收全部通过，"
                         "7 次模型启动，耗时约 14.6 秒"],
        ["三模态归档推理（河南）", "Excel、HTML、扫描 PDF 共 17 处定位进入四节点计划，Nexent 实跑 15 项验收全部通过，"
                         "7 次模型启动，耗时约 13.2 秒"],
        ["自由提问问答", "Nexent 智能体通过检索、核验两个 MCP 工具作答；界面实测跨文档收费问题 4.07 秒给出带两处引用的答复，"
                   "询问“现在的收费”时按规则停止作答"],
        ["本体半自动构建（北京）", "从归档办事指南生成新类型和新关系，人工确认后在本地执行 6 个虚构场景，随后回退"],
        ["自主更新（河南）", "心跳在自然到期时发现新归档资料，整理与核验子 Agent 生成 1 条待审候选，正式版本未被改动"],
        ["运行保障", "主控进程退出后由计划任务约 4 秒恢复监听；官方来源观察完成首次采集并形成待审修订"],
        ["工程交付", "公开仓库以 MIT 许可开源；工程包解压到新目录后可自行重建，包内 392 项离线测试通过"],
    ], [4.8, 9.8])

    w.heading("1.5 应用价值与资产激活路径", 2)
    w.para("本项目把沉睡的公开资料逐步激活为可用的决策依据，路径为：原件登记与定位 → 检索单元与本体候选 → "
           "人工确认的可执行计划 → 带出处的答复与更新候选。在这条路径上，三类用户分别获益：")
    w.bullets([
        ("群众：", "直接得到带原文出处的答复；资料不足时会被明确告知原因和缺少的依据，不会拿到可能错误的结论。"),
        ("工作人员：", "答复绑定资料版本和定位，便于复核与交接；新旧政策的差异以待审候选形式呈现，审核后再生效。"),
        ("维护者：", "新增事项或地区时复用 Skill 模板和 MCP 工具，只需配置资料来源和推理计划，不必从零搭建智能体。"),
    ])

    w.heading("1.6 当前边界", 2)
    w.para("本文所述结果均来自冻结的公开历史资料。系统答复统一标注“历史归档解释、未核定现行效力”，不构成政策发布。"
           "业务效果的独立评测（30 题配对比较）和低资源构建效率对照正在进行，结果将在效果验证部分给出；"
           "扫描件目前使用已有的署名 AI 转录，人工复核正在进行。")


def chapter_two(w):
    w.heading("2 整体方案设计", 1)
    w.heading("2.1 设计原则", 2)
    w.bullets([
        ("来源可溯：", "每条结论绑定原件哈希、修订和定位，引文必须逐字出自原文。"),
        ("人工把关：", "新的本体结构和新的资料版本先作为待审候选，经人工确认才生效；结构确认不等于业务批准。"),
        ("确定性优先：", "检索、计划执行、数值与单位核对等可以计算的部分交给工具，模型不自行补全事实。"),
        ("失败即停：", "来源缺失、版本变化、问题超出范围或核验不通过时停止作答，并说明原因。"),
        ("最小权限与隐私：", "工具只接受限定输入，不接收身份证号等个人信息；凭据只保存在本机私有目录。"),
    ])

    w.heading("2.2 总体架构", 2)
    w.para("系统分为五个功能层和一组运行保障机制（图 2-1）。数据自下而上被识别、组织为知识，经工具执行和智能体编排，"
           "最终以带出处的答复呈现给用户；用户问题和更新任务则自上而下逐层下发。")
    w.figure("图1-系统总体架构.png", "图 2-1 系统总体架构", width=14.5)
    w.table("表 2-1 各层组成与实现", ["层次", "主要组成", "实现要点"], [
        ["数据资产层", "官方 HTML、Excel 收费目录、扫描 PDF", "采集前检查 robots 规则，核对原件 SHA-256；按页面选择器提取段落和条文，"
                                           "Excel 只读字面值、不执行公式，扫描页保留页码与左右半页定位"],
        ["知识与本体层", "本体种子、候选、版本库", "种子含 7 类实体（事项、条件、材料、部门、程序、规范条款、资产）和 6 类关系"
                                       "（需要、适用于、办理部门、包含步骤、对齐、证据支持）；候选经人工确认后成为不可变版本，可回退"],
        ["工具层", "MCP 服务", "基于 FastMCP 的 SSE 服务，使用专用 Bearer 令牌、Host 与 Origin 白名单，"
                           "启动时记录代码指纹，代码变更后拒绝继续服务"],
        ["编排层", "Nexent 主控与子 Agent、Skill", "主控按 Skill 规定的顺序调度子 Agent，每个角色只能使用指定工具，"
                                          "并设有步数上限"],
        ["交互层", "Nexent 对话与调试界面", "群众或工作人员用自然语言提问，答复包含结论、逐条引用、未知项和状态说明"],
        ["运行保障", "心跳、计划任务、来源观察", "主控心跳与计划任务恢复，官方来源到期采集，暂停门与人工确认门，"
                                     "调用预算与回执审计"],
    ], [2.6, 4.2, 7.8])

    w.heading("2.3 核心执行流程", 2)
    w.heading("2.3.1 检索-推理双驱动查询流程", 3)
    w.para("查询流程如图 2-2 所示。主控 Agent 读取 Skill 后先做范围检查：问题含个人信息、涉及试点范围以外的地区，"
           "或询问“现在、现行”的规定时，直接停止并说明原因。范围内的问题进入检索，工具取回前 5 条证据；本体推理链还会检查来源是否暂停或已变化。"
           "推理阶段，固定业务情境执行已确认的本体依赖计划；自由提问则由模型在规则约束下组织答复。"
           "最后，核验工具逐条检查引文是否逐字出自检索到的原文、答复中的数字和天数单位是否被引用，不通过即改为停止作答。")
    w.figure("图2-检索推理双驱动查询流程.png", "图 2-2 检索-推理双驱动查询流程")
    w.para("以南京试点为例，系统需要同时读取两份官方问答：周期问答给出省内 20 个工作日、省外户籍增加 5 个工作日，"
           "旧证问答说明“未超期的旧身份证可以继续使用”。本体计划依次执行三步：计算办理周期、取旧证有效区间的交集、比较二者覆盖关系，"
           "在虚构情境（省外户籍、旧证剩余 22 个工作日）下得出“不能确认旧证可覆盖整个办理周期”的结论。"
           "缺少任一份文件时，结论自动降为未知。")
    w.heading("2.3.2 知识更新与本体演化流程", 3)
    w.para("更新流程如图 2-3 所示。主控心跳和官方来源观察按到期时间采集资料，发现新原件后记录哈希与修订；"
           "整理子 Agent 分别读取新旧资料的完整视图并提交结构化差异，核验子 Agent 独立复核证据；结果以待审候选形式保存，"
           "不继承旧批准，也不自动发布。人工确认后，新版本才被激活；版本不可变，可随时回退。"
           "本体结构的变化（类型、关系、执行语义）与资料内容的变化（修订、差异候选）分开管理。")
    w.figure("图3-知识更新与本体演化流程.png", "图 2-3 知识更新与本体演化流程")

    w.heading("2.4 智能体设计", 2)
    w.para("系统在 Nexent 上配置三组智能体（表 2-2）。所有智能体使用 DeepSeek 官方模型 deepseek-flash，温度 0，"
           "关闭思考模式，每步输出不超过 1024 tokens；演示期间均为禁用、未发布的草稿，由受控接口或配置页调试对话运行。")
    w.table("表 2-2 智能体与职责", ["智能体组", "角色", "可用工具", "关键约束"], [
        ["自由提问", "问答助手（主）", "检索、核验", "最多 6 步；只依据检索证据作答；核验不通过即停止"],
        ["本体推理", "主控", "开启任务、读取状态、完成交付", "最多 6 步；只有持久化的核验阶段通过后才能交付"],
        ["本体推理", "检索子 Agent", "取回绑定来源的完整原文", "最多 3 步；每个任务只调用一次"],
        ["本体推理", "核验子 Agent", "执行已确认的本体依赖计划", "最多 3 步；使用工具返回的实际参数和依赖轨迹"],
        ["知识更新", "主控", "读取阶段状态、完成", "先调用整理、后调用核验；完成门通过才返回"],
        ["知识更新", "整理子 Agent", "读取新旧视图、提交差异", "最多 4 步；新旧资料各读一次完整视图"],
        ["知识更新", "核验子 Agent", "读取新旧视图、提交复核", "最多 4 步；独立核对证据与候选"],
    ], [2.4, 3.0, 4.0, 5.2])

    w.heading("2.5 MCP 工具与 Skill 分层编排", 2)
    w.para("工具按职责拆分，每个工具只做一件事（表 2-3）。每次调用都会重新检查服务代码指纹；"
           "本体推理工具还会复查来源门和已确认的本体版本。")
    w.table("表 2-3 MCP 服务与工具", ["服务", "工具", "作用"], [
        ["本体推理服务", "open_query", "按固定情境开启任务，返回任务编号"],
        ["本体推理服务", "retrieve", "取回本体绑定来源的完整原文选段"],
        ["本体推理服务", "verify", "执行已确认的本体依赖计划，记录逐跳依据"],
        ["本体推理服务", "state / finalize", "读取持久阶段并复查来源与版本；重新执行并原样返回持久答复"],
        ["自由提问服务", "frozen_retrieve", "范围检查，返回前 5 条证据和作答规则"],
        ["自由提问服务", "frozen_verify", "核对引文、数字和天数单位，不通过即改为停止"],
        ["知识更新服务", "get / submit / review / state / finish", "读取新旧视图，提交差异与复核，推进阶段并完成"],
    ], [3.0, 5.2, 6.4])
    w.para("Skill 采用分层编排（表 2-4）。编排层和协议层已沉淀为参数化模板：新部署只需填写情境、Agent 名称和证据说明，"
           "由渲染工具生成 Skill 与主控协议；执行层的来源适配和计划算子仍需按事项开发。")
    w.table("表 2-4 Skill 分层与复用方式", ["层次", "内容", "复用方式"], [
        ["平台层", "主控与检索、核验子 Agent，模型参数、步数上限", "统一接线脚本创建，profile 只换名称"],
        ["编排层", "读 Skill → 开任务 → 检索 → 核验 → 读状态 → 完成", "Skill 模板固定"],
        ["协议层", "可执行动作格式、子 Agent 调用示例、禁止重复与重试", "主控协议模板固定"],
        ["工具层", "5 个本体推理工具及其门禁检查", "共用服务，按 profile 分派"],
        ["执行层", "来源适配器、已确认依赖计划、领域算子", "按事项开发"],
    ], [2.4, 7.0, 5.2])

    w.heading("2.6 数据与安全设计", 2)
    w.bullets([
        ("数据范围：", "只使用公开的官方资料，标记为公开数据；演示使用虚构情境，不采集居民证件号或真实个人回答。"),
        ("凭据与网络：", "模型凭据通过环境变量或本机私有目录提供，不写入仓库；MCP 服务使用专用令牌；"
                    "云端请求沿用系统代理并保留 HTTPS 证书校验，本机 Nexent 与 MCP 服务直连。"),
        ("状态门：", "答复保持“未审核暂停、现行效力未知、未正式批准”的状态标注；新候选不自动发布，旧批准不被继承。"),
        ("预算与失败处理：", "调用前预留预算并设上限；模型、工具或来源出错时停止，不重试业务请求，失败记录完整保留。"),
    ])

    w.heading("2.7 部署与运行环境", 2)
    w.table("表 2-5 运行环境", ["组件", "版本与配置"], [
        ["智能体平台", "Nexent v2.7.0（Docker 部署）：配置服务 5010、运行服务 5014、Web 界面 3000"],
        ["语言模型", "DeepSeek 官方 API，deepseek-flash；温度 0，非思考模式"],
        ["MCP 服务", "本体推理 8785、自由提问 8786、知识更新 8784，均仅监听本机，通过 host.docker.internal 供容器访问"],
        ["运行环境", "Windows 11、Docker Desktop（WSL2）、Python 3.12，依赖版本锁定"],
        ["交付方式", "公开仓库（MIT 许可）；打包脚本生成工程包并核验散列，Agent 配置以平台导出 JSON 提供"],
    ], [3.2, 11.4])


def chapter_three(w):
    w.heading("3 知识图谱构建方案", 1)
    w.heading("3.1 构建思路", 2)
    w.para("政务知识图谱既要回答“依据是什么”，也要执行“条件是否满足”“区间能否覆盖”这类推理。本项目不追求一次性抽取全量图谱，"
           "而是在少量种子本体之上，围绕具体业务问题逐步扩展可执行的类型和关系。构建分为四步（图 3-1）：从官方原文的完整条款出发，"
           "用有界原文语法抽取候选；把类型、关系、条件、依赖计划及来源引文整理成候选摘要；由人工只确认结构；确认后的结构成为不可变版本，"
           "由执行器按版本执行。构建阶段不调用大模型，人工投入集中在一次结构确认上，适合标注稀缺的政务场景。")
    w.figure("图5-半自动构建流程.png", "图 3-1 低资源半自动本体构建流程")

    w.heading("3.2 本体模型设计", 2)
    w.heading("3.2.1 种子本体", 3)
    w.para("种子本体面向政务办理事项，定义 7 类实体（表 3-1）和 6 类关系（表 3-2），结构见图 3-2。"
           "种子是项目自建的工程结构，不等同于官方行业标准。")
    w.table("表 3-1 种子本体实体类型", ["实体", "中文名", "必填属性", "说明"], [
        ["Matter", "事项", "编号、名称、地区", "办理事项，业务关系的起点"],
        ["Condition", "适用条件", "编号、表达式、来源片段", "结果只取真、假、未知"],
        ["Material", "材料", "编号、名称、来源片段", "需交验或提交的材料"],
        ["Department", "部门", "编号、名称、地区", "受理或办理部门"],
        ["Procedure", "办理步骤", "编号、名称、来源片段", "办理程序中的环节"],
        ["StandardClause", "规范条款", "编号、引用、来源片段", "法规或规范性文件中的条款"],
        ["Asset", "原始资产", "编号、SHA-256、来源", "官方原件及其哈希"],
    ], [3.2, 2.4, 4.2, 4.8])
    w.table("表 3-2 种子本体关系类型", ["关系", "起点 → 终点", "含义"], [
        ["requires", "事项 → 材料", "办理需要的材料"],
        ["applies_when", "事项 → 适用条件", "事项适用的前提条件"],
        ["handled_by", "事项 → 部门", "受理或办理的部门"],
        ["has_step", "事项 → 办理步骤", "办理程序的组成"],
        ["aligned_to", "事项 → 规范条款", "事项与规范条款对齐"],
        ["supported_by", "规范条款 → 原始资产", "条款由哪份原件支持"],
    ], [3.4, 4.6, 6.6])
    w.heading("3.2.2 约束规则", 3)
    w.bullets([
        "每条业务关系必须绑定来源片段，没有出处的关系不能进入版本。",
        "跨地区合并必须显式复核，不同地区的参数不得互相套用。",
        "条件结果只取“真、假、未知”三值，信息不足时返回未知。",
        "候选分为待审、已接受、已拒绝三种状态；发布前必须经过人工审核。",
        "历史版本不可变，新内容只能以新版本追加。",
    ])
    w.heading("3.2.3 扩展机制", 3)
    w.para("新的业务问题常常需要种子中没有的类型或关系。扩展以“类型化依赖计划”表示：每个计划节点包含编号、算子、"
           "取自原文的参数、依赖的前序节点和来源键，执行器按依赖顺序求值，并记录每一跳的依据。"
           "目前已通过人工结构确认的扩展有三项，分别来自北京、南京和河南试点（图 3-2）。")
    w.figure("图4-本体结构.png", "图 3-2 种子本体与已确认的扩展")

    w.heading("3.3 低资源半自动构建方法", 2)
    w.heading("3.3.1 有界原文语法抽取", 3)
    w.para("抽取器针对一类资料的固定表述编写，例如北京办事指南中“符合以下条件可在网上办理”的标题及紧邻续段。"
           "它识别结构，从原文中捕获数值，生成候选类型、关系和条件实例。数值不写死在代码里：用改动了数字的独立文本测试，"
           "执行结果随原文数值变化，说明不是查表。抽取器只接受完整条款和续段；遇到额外条件、否定条件或结构变化即拒绝，交由人工处理。")
    w.heading("3.3.2 候选摘要与人工结构确认", 3)
    w.para("候选摘要包含新类型、新关系、条件或依赖计划，以及来源修订、原件哈希、逐段引文和预期的版本头。"
           "维护者只需判断结构是否正确，确认记录写明“仅确认结构，不构成业务批准”。写入时比对版本头（CAS），防止并发修改相互覆盖。")
    w.heading("3.3.3 版本化执行与回退", 3)
    w.para("版本以内容哈希标识，执行器只读取已确认版本中的类型、关系和参数。回退时版本头指回原基线，"
           "查询重新返回“不支持的关系”（UNSUPPORTED_RELATION），已确认的历史版本仍然保留，便于追溯。"
           "来源修订、原件或实现指纹任一发生变化，依赖它的旧结果都会被阻断。")
    w.heading("3.3.4 三地构建实例", 3)
    w.para("三个试点分别覆盖条件判断、跨文档区间比较和多模态归档组装三类推理（表 3-3）。")
    w.table("表 3-3 三地本体构建实例", ["试点", "来源与规模", "新增类型与关系", "构建与执行结果"], [
        ["北京", "1 份办事指南、2 段条款；1 个示例种子", "办理渠道（ServiceChannel）；事项 → available_via → 办理渠道；"
                                              "5 项合取条件，4 个阈值取自原文", "模型调用 0 次；人工确认后本地执行 6 个虚构场景，回退后返回“不支持的关系”"],
        ["南京", "2 份官方问答、5 段正文", "工作日区间比较（IntervalComparison）；事项 → compares_coverage；"
                                  "三节点依赖计划，2 个数值参数取自原文", "模型调用 0 次；本地 12 项检查通过，Nexent 原生运行 13 项验收通过"],
        ["河南", "Excel、HTML、扫描 PDF 共 3 份资料，17 处定位", "换领归档依据包（ArchivedRenewalEvidence）；"
                                                    "事项 → assembles_archival_basis；四节点依赖计划", "模型调用 0 次；Nexent 原生运行 15 项验收通过"],
    ], [1.4, 3.4, 5.6, 4.2])
    w.para("以北京为例，候选包含“本市户籍、年满 16 周岁、上次办证不超过 5 年、照片采集不超过 5 年、已采集指纹”5 项条件。"
           "虚构输入满足全部条件时结果为真；缺少指纹信息时为未知；恰好处于 5 年边界时也保持未知。"
           "结果为真只表示输入满足这份归档原文的结构条件，不表示本人可以办理。")
    w.heading("3.3.5 河南四节点依赖计划", 3)
    w.para("河南计划把三种格式的资料组织成一条可执行链（表 3-4）。表格只提供规范文号、不含金额，收费分支来自财政部通知正文；"
           "扫描件提供独立的历史服务条件，不是收费结论的前提。")
    w.table("表 3-4 河南四节点依赖计划", ["节点", "作用", "依赖", "来源", "输出"], [
        ["catalog_basis", "识别规范依据", "—", "Excel 收费目录", "文号：财综〔2007〕34号"],
        ["fee_branch", "历史换领收费分支", "catalog_basis", "财政部通知（HTML）", "换领 20 元；遗失补领、损坏换领 40 元；临时证 10 元（历史原文）"],
        ["service_archive", "历史服务措施", "—", "扫描 PDF 第 5、6 页", "办理周期由 40 日（自然日）压缩至 30 日；30 日未注明日类型"],
        ["archive_answer", "组装归档依据", "fee_branch、service_archive", "三类来源", "带 17 处定位的归档答复"],
    ], [2.8, 2.6, 2.8, 2.6, 3.8])

    w.heading("3.4 本体动态更新与演化", 2)
    w.para("演化分为两类：结构演化，即新增或修改类型、关系、执行语义；内容演化，即同一事项的资料出现新修订或后继文件。"
           "两类变化都先成为待审候选，流程见图 2-3。")
    w.para("内容演化以河南办理时限为例：心跳在自然到期时发现 2026 年的新政策解读，整理子 Agent 分别读取 2024 年旧措施和 2026 年新解读的"
           "完整视图并提出差异，核验子 Agent 独立复核后生成 1 条待审候选。候选保留“25 个自然日”“15 个工作日”以及中高考特殊时期"
           "急需用证考生“绿色通道 5 个工作日”等原文的单位和适用对象，不做异单位相减；精确施行日期未知时如实标注。"
           "旧依据继续保留，审核前正式版本不变。")
    w.para("结构演化沿用 3.3 的确认与版本机制。来源被暂停、撤权或修订时，依赖它的结果自动阻断，等待重新确认后才能恢复。")

    w.heading("3.5 与行业规范、标准体系对齐", 2)
    w.para("本体通过“规范条款”实体以及“对齐”“证据支持”两类关系连接规范依据。每条映射按“规范条款 → 本体字段或关系 → "
           "适用对象与地区 → 生效范围 → 例外 → 复核状态”记录，并绑定原件修订和定位，使任何一条业务结论都能回溯到具体条文。"
           "表 3-5 列出部分映射。")
    w.table("表 3-5 规范映射示例", ["规范依据", "映射的本体要素", "适用范围与例外", "复核状态"], [
        ["财综〔2007〕34号 第一项", "收费分支 → 规范依据", "换领每证 20 元；遗失补领、损坏换领 40 元；临时证 10 元。属于历史原文分支，不是个案应缴金额",
         "待人工复核"],
        ["财税〔2018〕37号 第一条", "首次申领 → 停征规则", "2018 年 4 月 1 日起停征首次申领工本费；不扩大到换领、补领", "已做有范围的 AI 复核"],
        ["豫公通〔2022〕77号 第一百五十九条", "申领/换领/补领 → 需要 → 材料", "申领交验户口簿；换领交验居民身份证；补领交验有效身份证件之一",
         "待人工复核；已与另一份官方公开副本核对一致"],
        ["豫公通〔2022〕77号 第二百三十六条", "规范 → 施行日期", "条文载明 2022 年 9 月 1 日起施行，与文件签发日期分开记录", "待人工复核"],
        ["河南 2026 年便民措施 第一项", "服务时限 → 取值", "本省户籍一般 15 个工作日；5 个工作日只适用于中高考特殊时期急需用证考生",
         "已做有范围的 AI 复核"],
    ], [3.4, 3.4, 5.4, 2.4])
    w.para("目前已完成 10 条映射的来源与定位核对，逐条人工复核正在由队员进行。本项目不声称已通过国家或行业标准（如 GB/T 系列）"
           "的符合性认证。复核通过的映射将登记为“对齐”关系；规范发生变化时，据此提示受影响的查询路径。")

    w.heading("3.6 低资源构建效率", 2)
    w.para("三地构建均未使用标注数据，构建阶段模型调用均为 0 次。北京只用了 1 个示例种子，人工投入是对一份候选摘要做结构确认。")
    w.para("为量化效率，项目设计了对照实验：选取 6 个预先封闭的映射单元，覆盖适用条件、分支材料、规范效力、收费数值与例外、"
           "版本变化和扫描件跨页条款。方法 A 由人从原文手工整理，方法 B 先由系统生成候选、再由人复核；两位队员交叉分配单元，"
           "记录阅读、整理、复核的主动用时，以及直接采纳率、可用率和漏项数；抽取规则的编写成本单独记录。实验结果将在效果验证部分给出。")


def chapter_four(w):
    w.heading("4 智能体构建方案", 1)
    w.heading("4.1 智能体总体设计", 2)
    w.todo("李嘉杰", "说明三组智能体（本体推理、知识更新、自由提问）的定位、分工与协作关系，以及为什么采用“主控 + 子 Agent”结构；"
                    "可直接引用表 2-2，不重复列表。", "docs/competition/当前Nexent配置与证据说明-2026-10-21.md 第一至三节。")
    w.heading("4.2 本体推理智能体组", 2)
    w.heading("4.2.1 角色与提示词设计", 3)
    w.todo("李嘉杰", "主控、检索、核验三个角色的职责提示词、约束提示词和示例动作；说明为什么禁止导入模块、禁止改写工具结果。",
           "scripts/provision_ontology_crossdoc_native.py 的 PROTOCOL 与 draft_body；config/skills/govkb-ontology-crossdoc-preview。")
    w.heading("4.2.2 执行流程与完成门", 3)
    w.todo("李嘉杰", "按实际工具顺序写出一次完整运行：读 Skill → open_query → 检索子 Agent → 核验子 Agent → state → finalize；"
                    "说明“只有持久化的核验阶段通过才能交付”的完成门。",
           "南京原生统一验收-2026-10-21.md、河南三模态原生验收-2026-10-22.md。")
    w.heading("4.3 知识更新智能体组", 2)
    w.todo("李嘉杰", "主控、整理、核验三个角色；新旧资料各读一次完整视图；差异候选的字段；为什么候选不继承旧批准。",
           "config/skills/govkb-successor-update；主控心跳与自主更新修复验收-2026-10-21.md。")
    w.heading("4.4 自由提问问答助手", 2)
    w.todo("李嘉杰", "单智能体 + 检索、核验两个 MCP 工具的设计；答复格式（结论、逐条引用、未知项、状态说明）；停止规则。",
           "自由提问入口与评测运行说明-2026-10-30.md；scripts/provision_frozen_qa_native.py 的 PROMPT。")
    w.heading("4.5 Skill 设计与模板化", 2)
    w.todo("陈子炫", "Skill 的结构与分层（引用表 2-4），模板渲染流程，河南 Skill 逐字复现的验证，新 profile 部署步骤与计时结果。",
           "Skill工作流模板说明-2026-10-30.md；李嘉杰的模板试用计时记录。")
    w.heading("4.6 MCP 工具实现与关键代码", 2)
    w.todo("陈子炫", "选 2–3 段关键代码解读：工具的指纹检查与失败即停、核验函数（引文逐字、数字与单位）、依赖计划执行器。"
                    "每段代码不超过 20 行，配一句说明。",
           "src/gov_mcp/frozen_qa.py 的 verify；src/gov_mcp/ontology_crossdoc_preview_server.py 的 execute；crossdoc_reasoning.py。")
    w.heading("4.7 输出契约与容错", 2)
    w.todo("陈子炫", "严格 JSON 输出契约、工具结果原样交付、模型或工具出错即停、不重试业务请求、预算预留与调用上限。",
           "JSON稳定性修复-2026-10-21.md；严格JSON输出修复-2026-10-19.md。")


def chapter_five(w):
    w.heading("5 原始数据说明", 1)
    w.heading("5.1 数据来源与选取原则", 2)
    w.todo("程千城", "说明只选用政府官网公开资料的原因；为什么以居民身份证业务为试点；为什么选河南作为主要评测范围，南京、北京作为补充；"
                    "资料截至时间与冻结方式。", "evaluation/independent-20261030/scope.md；真实资产与规范对齐表-2026-10-21.md。")
    w.heading("5.2 冻结资料清单", 2)
    w.para("河南试点的评测与问答使用以下 7 份官方原件（表 5-1）；南京、北京试点资料见表 5-2。")
    w.table("表 5-1 河南冻结资料清单", ["编号", "资料", "格式", "发布或转载机构", "主要用途"], [
        ["A2", "河南省执行的全国性及中央部门和单位行政事业性收费目录清单", "Excel", "河南省财政厅", "收费项目与依据文号"],
        ["A3", "河南公安 2021 年便民利企措施", "扫描 PDF", "河南省公安厅", "历史省内通办、办理周期"],
        ["A4", "关于进一步规范居民身份证收费等有关问题的通知（财综〔2007〕34号）", "HTML", "财政部", "换领、补领、临时证收费分支"],
        ["A5", "停征免征和调整部分行政事业性收费有关政策（财税〔2018〕37号）", "HTML", "财政部", "首次申领停征工本费"],
        ["A6", "河南公安 2026 年便民利民措施政策解读", "HTML", "河南省公安厅", "办理时限新规定"],
        ["A7", "河南省户口居民身份证管理工作规范（试行）（豫公通〔2022〕77号）", "HTML", "河南省公安厅", "申领、换领、补领的材料与程序"],
        ["B1", "河南 2024 年便民措施（转载）", "HTML", "济源市公安局", "办理时限旧规定"],
    ], [1.2, 5.6, 1.6, 2.6, 3.6])
    w.table("表 5-2 其他试点资料", ["试点", "资料", "格式", "用途"], [
        ["南京", "南京市政府问答：身份证到期换证周期（2026-01-05）", "HTML", "省内、省外户籍的换发周期"],
        ["南京", "南京市政府问答：换领期间旧证能否使用（2025-12-09）", "HTML", "未超期旧证的使用条件"],
        ["北京", "北京市居民身份证换领办事指南", "HTML", "网上办理条件"],
    ], [1.6, 7.0, 1.6, 4.4])
    w.todo("程千城", "核对表 5-1、5-2 的资料名称与机构；补充每份资料的发布日期（页面日期，不等于施行日期）和原件 SHA-256 前 16 位。")
    w.heading("5.3 数据规模与分布", 2)
    w.table("表 5-3 数据规模", ["口径", "规模", "说明"], [
        ["资料目录（10 月 18 日快照）", "95 个站点、79 份文档、186 个修订、1476 个片段", "多地区资料登记，含待审与排除状态"],
        ["河南冻结语料", "7 份原件、581 个检索单元", "Excel 行 258、规范条文 236、网页段落 85、扫描转录 2"],
        ["河南三模态推理链", "3 份资料、17 处定位", "Excel 单元格 10、HTML 条款 5、扫描转录 2"],
    ], [4.2, 5.0, 5.4])
    w.todo("程千城", "补充各地区、各格式的文档数分布（可做一张柱状图）。", "data/knowledge/candidate-catalog.json；evidence/frozen-qa-20261030/corpus-manifest.json。")
    w.heading("5.4 数据权利与使用边界", 2)
    w.todo("程千城", "说明资料均为公开官方资料、只用于评测与演示；原件再分发许可尚未逐一核定，因此公开仓库只提供网址与散列；"
                    "不采集个人信息，演示使用虚构情境。", "复现与交付说明-2026-10-21.md 第 2 节。")


def chapter_six(w):
    w.heading("6 数据处理说明", 1)
    w.heading("6.1 处理流程总览", 2)
    w.todo("程千城", "用一张流程图或一段话概括：采集 → 原件核验 → 解析与定位 → 检索单元 → 本体候选 → 人工复核 → 版本。"
                    "可复用图 2-3 的风格，由陈子炫出图。")
    w.heading("6.2 采集与原件核验", 2)
    w.todo("程千城", "采集前检查 robots 规则；每份原件计算 SHA-256 并与登记值核对，不一致即停止；网站无法访问时使用历史保存的同一原件"
                    "（如济源公安网站 TLS 握手失败，改用更新链保存的原件，散列一致）。", "scripts/build_frozen_corpus.py；evidence/frozen-qa-20261030/corpus-manifest.json。")
    w.heading("6.3 多模态解析与定位", 2)
    w.table("表 6-1 各格式的解析与定位方式", ["格式", "解析方式", "定位粒度", "注意事项"], [
        ["HTML", "按页面正文选择器提取段落；规范文件按“第 X 条”合并续行", "正文段落 N / 第 X 条", "保留完整条款和续段，不截断"],
        ["Excel", "只读字面值，不执行公式；按行组织", "工作表!第 N 行、单元格", "公式单元格跳过并计数"],
        ["扫描 PDF", "按页渲染后人工或 AI 转录", "PDF 页码、左右半页", "转录与原件分开保存，需人工复核"],
    ], [2.0, 5.0, 3.4, 4.2])
    w.todo("程千城", "按表 6-1 逐项展开，各配一个真实例子（如第一百五十九条、收费目录第 36 行、扫描件第 5 页右半页）。")
    w.heading("6.4 检索单元构建", 2)
    w.todo("陈子炫", "检索单元的字段（来源、定位、原文、散列、证据类型）；字面二元组检索与前 5 条证据；与目录登记条款散列的复核（22 条一致）。",
           "src/gov_mcp/frozen_qa.py；scripts/build_frozen_corpus.py。")
    w.heading("6.5 质量控制与人工复核", 2)
    w.todo("程千城", "扫描件真人复核结果（逐处一致或差异）；规范条款映射复核结果；被排除资料及原因。",
           "docs/competition/扫描件真人复核-程千城-*.md；规范条款映射复核-*.md。")
    w.heading("6.6 隐私与安全处理", 2)
    w.todo("程千城", "不接收身份证号等个人信息（含 18 位号码的问题直接停止）；凭据与运行数据只放本机私有目录；公开证据只含数量与散列。")


def chapter_seven(w):
    w.heading("7 Nexent 智能体配置说明", 1)
    w.heading("7.1 平台与模型信息", 2)
    w.table("表 7-1 平台与模型配置", ["项目", "配置"], [
        ["智能体平台", "Nexent v2.7.0，Docker 部署于本机"],
        ["语言模型", "DeepSeek 官方 API，deepseek-flash；所有智能体使用同一模型"],
        ["解码参数", "温度 0，关闭思考模式，每步输出上限 1024 tokens"],
        ["步数上限", "主控与问答助手 6 步；检索、核验子 Agent 3 步；更新组整理、核验子 Agent 4 步"],
        ["发布状态", "全部为禁用、未发布的草稿；通过受控接口或配置页调试对话运行"],
    ], [3.2, 11.4])
    w.heading("7.2 工具信息", 2)
    w.para("各 MCP 服务的工具及作用见表 2-3。服务以 SSE 方式注册到 Nexent，使用专用 Bearer 令牌鉴权，容器通过 "
           "host.docker.internal 访问本机服务。")
    w.todo("李嘉杰", "补一张 Nexent“MCP 配置”页面截图，并列出每个服务的注册名、地址和工具数量。", "截图 S02。")
    w.heading("7.3 知识库信息", 2)
    w.para("当前知识载体是由 MCP 服务管理的冻结语料和来源绑定包：河南 7 份原件整理为 581 个检索单元，南京两份问答、"
           "河南三模态资料分别绑定在对应的本体推理链中。")
    w.todo("陈子炫", "确定嵌入模型后，把冻结语料导入 Nexent 原生知识库，补充知识库名称、文件列表、切片方式与截图；"
                    "未导入前如实说明当前方式。", "scripts/export_frozen_knowledge.py（已导出 7 个 Markdown 文件）；截图 S08。")
    w.heading("7.4 调用关系图", 2)
    w.todo("李嘉杰", "画一张调用关系图：用户 → 主控 → 子 Agent → MCP 工具 → 本体/语料，标出每一步调用的工具名和顺序；"
                    "放在 docs/competition/assets/design/ 下，图题为“图 7-1 智能体调用关系”。", "表 2-2、表 2-3；Nexent 调试页的调用过程截图 S04c。")
    w.heading("7.5 调试迭代经验", 2)
    w.para("开发过程中遇到的主要问题及处理见表 7-2，所有失败运行都保留了记录。")
    w.table("表 7-2 主要调试问题与处理", ["问题", "原因", "处理"], [
        ["子 Agent 读到的正文被截断", "长回执在平台观察中被压缩，中间正文丢失", "新旧资料拆成两份完整视图分别读取，超长即阻断"],
        ["整理未提交即达步数上限，核验仍被启动", "缺少阶段门控", "主控按持久阶段调度，前一阶段未完成不启动下一阶段"],
        ["子 Agent 未在同一动作中结束、把字符串当字典", "模型对工具返回值的处理不稳定", "要求工具调用与结束在同一动作，工具结果原样交付，固定状态标记"],
        ["沙箱中导入模块失败", "Nexent 代码沙箱限制导入", "提示词禁止导入；核验工具改为接收结构化字段"],
        ["验收脚本把 true 与 1、25 与 25.0 判为相等", "回执比较使用普通相等，不区分 JSON 类型", "改为递归按类型严格比较，并补变异回放检查"],
        ["Docker 启动报缺少注册表键", "此前在其他应用环境中安装，数据与设置被重定向", "按原版本修复安装，复制原数据盘并恢复设置，17 个容器全部恢复"],
    ], [4.4, 4.8, 5.4])
    w.todo("李嘉杰", "核对表 7-2，并补充每个问题前后各一次运行的对比（例如河南更新链 A1–A6 六轮，A6 成功）。",
           "当前Nexent配置与证据说明-2026-10-21.md 第四节；JSON稳定性核查-2026-10-21.md；接续记录第 42 节。")
    w.heading("7.6 示例问答截图", 2)
    w.para("图 7-2 为自由提问助手对跨文档收费问题的答复，图 7-3 为询问“现在的收费”时按规则停止作答。")
    w.figure("S04a-跨文档问答结论与引用-20261030.jpg", "图 7-2 跨文档问答：结论与逐条引用", width=14, folder=SCREENSHOTS)
    w.figure("S05a-现行效力问题停止作答-20261030.jpg", "图 7-3 询问现行收费时停止作答并列出归档原文", width=14, folder=SCREENSHOTS)
    w.todo("陈子炫", "补充配置页截图（S01a/S01b）和调用过程截图（S04c）；南京、河南推理链截图（S06/S07）需要新运行或注明原运行日期；"
                    "正式交稿前用 Win+Shift+S 重截高清版。", "docs/competition/assets/screenshots/README.md。")
    w.heading("7.7 同步提交文件说明", 2)
    w.table("表 7-3 同步提交文件", ["类别", "文件", "说明"], [
        ["Agent JSON", "config/exports/ontology-crossdoc-native-20261021/agent.json；config/exports/henan-multimodal-native-20261022/agent.json",
         "Nexent 平台导出的智能体配置，原始导出 ZIP 在 artifacts/ 下"],
        ["Skill", "config/skills/ 下的在用 Skill；config/skill-templates/ 模板与规格", "含分层模板与渲染工具"],
        ["MCP 文件", "src/gov_mcp/ontology_crossdoc_preview_server.py；src/gov_mcp/frozen_qa_server.py 及依赖模块", "附启动命令与端口说明"],
        ["知识库", "export_frozen_knowledge.py 导出的 7 个 Markdown 文件", "原件再分发许可未核定，提交前须确认"],
        ["工程包", "build_competition_package.py 生成的 ZIP 与散列清单", "MIT 许可；可在新目录自行重建"],
    ], [2.4, 7.2, 5.0])
    w.todo("陈子炫", "按官方上传要求确认文件格式与打包方式后定稿本表。")


def chapter_eight(w):
    w.heading("8 效果验证与总结", 1)
    w.heading("8.1 验证方法", 2)
    w.para("效果验证分两部分：一是业务效果，用 30 道独立编写的问题（20 道应答复、10 道应停止）对比普通检索与本方案；"
           "二是低资源构建效率，用 6 个映射单元对比纯人工整理与“系统候选 + 人工复核”。两种方案使用同一批资料、同一模型和同一检索，"
           "答复去掉方法名并随机排序后盲评。")
    w.todo("陈子炫", "补充出题人、冻结时间、盲评人和独立性声明；说明 30 题是项目的评测设计，不是官方规定。",
           "独立评测执行方案-2026-10-21.md；evaluation/independent-20261030/。")
    w.heading("8.2 业务效果评测结果", 2)
    w.table("表 8-1 业务效果对比（待测）", ["指标", "普通检索", "本方案", "说明"], [
        ["严格整题正确率", "待测", "待测", "结果类型、关键断言、范围边界全部正确的题数占比"],
        ["断言支持率", "待测", "待测", "有原文支持的关键断言占比"],
        ["前 5 条证据召回率", "待测", "待测", "仅统计应答复题"],
        ["正确停止率", "待测", "待测", "应停止题中实际停止的比例"],
        ["误停止率", "待测", "待测", "应答复题中错误停止的比例"],
        ["平均耗时", "待测", "待测", "从提问到答复的时间"],
        ["平均输入 tokens", "待测", "待测", "每题模型输入用量"],
    ], [3.6, 2.2, 2.2, 6.6])
    w.todo("陈子炫", "先跑 8 题先导，再跑 30 题；填写表 8-1，并逐类分析错误（单位、分支、地区、版本）。")
    w.heading("8.3 低资源构建效率结果", 2)
    w.table("表 8-2 低资源构建对照（待测）", ["单元", "方法 A 主动用时", "方法 B 主动用时", "B 直接采纳率", "B 可用率", "漏项（A / B）"], [
        ["U1 申领条件", "待测", "待测", "待测", "待测", "待测"],
        ["U2 分支材料", "待测", "待测", "待测", "待测", "待测"],
        ["U3 规范效力", "待测", "待测", "待测", "待测", "待测"],
        ["U4 收费数值与例外", "待测", "待测", "待测", "待测", "待测"],
        ["U5 版本变化", "待测", "待测", "待测", "待测", "待测"],
        ["U6 扫描件跨页条款", "待测", "待测", "待测", "待测", "待测"],
    ], [3.2, 2.4, 2.4, 2.2, 2.0, 2.4])
    w.todo("程千城、李嘉杰", "按 evaluation/lowresource-20261030/ 的记录表交叉完成方法 A、B 并汇总；U4–U6 在开发中用过，需注明方法 B 可能偏乐观。")
    w.heading("8.4 工程验收汇总", 2)
    w.para("各能力的工程验收结果见表 1-1。这些结果证明系统在冻结资料上的实现与边界，不等同于业务准确率。")
    w.heading("8.5 局限与后续工作", 2)
    w.todo("陈子炫", "如实列出：结果基于历史归档资料；扫描件依赖转录；新事项的执行层仍需开发；多行业迁移（总决赛附加分方向）。")
    w.heading("8.6 总结", 2)
    w.todo("陈子炫", "用一段话回扣四个评分维度：编排与 MCP/Skill 集成、资产激活业务价值、低资源本体构建、交互体验与能力沉淀。")


def main():
    doc = Document()
    configure(doc)
    writer = Writer(doc)
    title_block(writer)
    chapter_one(writer)
    # No forced chapter break: pagination is redone when pasted into the official template.
    chapter_two(writer)
    chapter_three(writer)
    for chapter in (chapter_four, chapter_five, chapter_six, chapter_seven, chapter_eight):
        chapter(writer)
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
