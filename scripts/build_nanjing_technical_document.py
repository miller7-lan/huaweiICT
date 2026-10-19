"""Build the technical handoff from frozen public evidence; use bundled document Python."""
import hashlib
import json
from pathlib import Path
import subprocess

from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT

ROOT = Path(__file__).resolve().parents[1]


def read(relative): return json.loads((ROOT / relative).read_text(encoding="utf-8-sig"))


def main():
    reports = sorted(p for p in (ROOT / "evidence/real-matter-closure/native-dataflow-v4").glob("20*.json") if "trace" not in p.stem)
    if not reports: raise RuntimeError("Wait for the bounded native run's terminal public receipt")
    native = read(reports[-1].relative_to(ROOT))
    pack = read("data/real-matter-closure/nanjing/update-snapshots-v1/registered-four/rules-and-sources.json")
    update = read("evidence/real-matter-closure/update-cycle/stage-04-update-live.json")
    clean = read("evidence/real-matter-closure/project-delivery/extracted-http-reproduction.json")
    ocr = read("evidence/real-matter-closure/ocr-handoff/interface-receipt.json")
    download = read("evidence/real-matter-closure/project-delivery/github-download.json")
    assert pack["knowledge_published"] is False and pack["human_review"]["status"] == "pending"
    assert clean["status"] == ocr["status"] == "PASSED" and download["matches"]
    assert len(pack["sources"]) == 4 and sum(len(s["segments"]) for s in pack["sources"]) == 142 and len(pack["rules"]) == 16
    assert len(update["checks"]) == len(clean["checks"]) == 19
    assert all(c["passed"] for c in update["checks"] + clean["checks"])
    assert native["native_run_requests"] == 1 and native["campaign"]["closed"] and native["campaign"]["active"] is None
    commit = subprocess.check_output(["git", "rev-parse", "HEAD"], cwd=ROOT, text=True).strip()
    status = "通过" if native.get("passed") else "未通过"
    pages = [
      ("南京居民身份证期满换领技术说明", [
        ("p", "本说明供项目队员和指导教师核查南京真实事项的工程实现与交付。当前已具备来源绑定、候选本体与跨文档推导、更新工程发布、群众与工作人员交接及解压项目复现。规则、关系、OCR和答复继续保持候选，正式业务发布关闭；独立业务评测没有结果。"),
        ("p", "工程记录整理日期为2026年10月19日，由AI辅助编写，未代签任何人工业务复核。本文采用项目技术说明格式，尚未套用官方赛事Word模板。"),
        ("h", "事项与适用范围"),
        ("p", "主分支为南京市居民身份证期满换领，成年人、具有完全民事行为能力、本人办理、江苏户籍。旧证生命周期推导另须声明旧证仍有效。未知、未成年人、代办、省外户籍、遗失损坏、历史日期及证据不足不能自动套用主分支。所有演示条件均为虚构声明，不收集姓名、证件号、联系方式或实际申请附件。"),
        ("table", (["项目", "已验证范围", "状态"], [
          ["真实资料", "4份HTML 142段正文及4页扫描公报", "候选原件"],
          ["规则核查", "16条来源绑定候选规则", "人工待核"],
          ["本体与多跳", "4条替代路径支持1个候选结论", "工程可执行"],
          ["更新与交接", "实际HTTP 19项检查", "工程通过"],
          ["解压复现", "新Python环境 双端HTTP 19项", "工程通过"],
          ["本地Qwen v4", "实际主子Agent轨迹", status],
          ["独立评测", "尚无独立冻结案例和评分", "待完成"]])),
        ("p", "工程检查数、来源段落数、替代证据路径数均不是独立业务案例数。本文不据此报告政策准确率、人工节时或模型学习效果。")]),
      ("官方原件与OCR候选", [
        ("p", "完整原件、哈希、段落定位、页面日期观察和抓取回执保存在来源包与更新快照中。回答引用必须同时携带原文位置和适用范围。公布日、施行日和抓取日分开记录，未知日期保持未知。"),
        ("table", (["来源", "正文量", "用途"], [
          [s["source_id"], str(len(s["segments"])), s["title"]] for s in pack["sources"]])),
        ("h", "来源与许可边界"),
        ("p", "本轮四份HTML均实际返回200，既有正文未观察到政策变化。新增法律来源属于受控登记范围增加，不能称为官网发布新政策。robots观察和请求回执保留；访问停止、正文或元数据变化须进入复核，不自动恢复。"),
        ("h", "OCR交接职责"),
        ("p", "OCR由队友负责实现，本轮只调用仓库已有Windows OCR接口一次，上游两份实现哈希未变，无新模型或识别引擎。真实四页返回通过现有契约校验，保存原始识别与可检索派生文本到隔离候选库；复核记录、发布版本和发布头均为0。"),
        ("p", "AI查看印刷页651至654。第1页混有修改决定与重公布法律，第4页法律末尾下方另起草案说明，不能将整个PDF等同法条。抽查可见缺字、错字及双栏顺序问题，未用模板修补。第2页的交回原证和六十日等锚点仅为观察，不代表逐字识别正确率。OCR仍待人工核对与文类切分，尚未用于发布问答证据。"),
        ("p", "来源索引与完整定位见 docs/南京真实事项来源与规则复核包.md；识别原始返回、请求和视觉边界见 data/real-matter-closure/nanjing/ocr-handoff/ 与 evidence/real-matter-closure/ocr-handoff/。")]),
      ("低资源本体与跨文档推导", [
        ("p", "当前方案使用少量种子关系及来源绑定模板提出候选，保留引用、范围、日期和人工待核标记。程序执行类型和引用校验，模型只负责宿主指导的工具编排。此实现不等同模型自主发现本体，也未测量低资源标注效率优势。"),
        ("h", "演化对照"),
        ("table", (["阶段", "来源与本体", "旧证生命周期查询"], [
          ["基线", "3来源49段 种子本体", "停止完整推导"],
          ["仅增来源", "4来源142段 种子本体", "仍停止完整推导"],
          ["关系演化", "4来源142段 新关系及组合", "4路径 1候选结论"]])),
        ("p", "候选关系连接办理期间使用条件、办理和领取阶段顺序，以及法律中的交回原证义务。阶段顺序是显式工程流程种子，不能冒充法律原文。地方来源有两处相关原文，法律HTML有修改决定与重公布条文两处对应内容，组合为四条替代证明，不是四个独立问题。"),
        ("code", "地方原文及条件 → 旧证在办理阶段可用\n流程种子 → 办理阶段先于领取阶段\n法律原文 → 领取新证时交回原证\n带类型的组合 → 条件限定的候选生命周期说明"),
        ("h", "可追溯与停止"),
        ("p", "每条证据绑定source_id、segment_id、原文SHA256和定位；持久版本封存原件及候选依赖。缺少关键资料、引用覆盖不足、版本切换竞争、来源暂停或新鲜度失效时停止。旧任务继续固定旧版，但历史答复也叠加当前来源控制，不能通过回退绕过暂停。"),
        ("p", "更新工程版本及实际对照见 evidence/real-matter-closure/update-cycle/；确定性重算与开发用例不代替独立业务评测或人工核定。")]),
      ("Nexent编排与执行约束", [
        ("p", "实际平台使用已登记的Qwen3 8B主控和两个Qwen3 4B子Agent，连接本机Ollama。主控绑定项目Skill，子Agent分别仅绑定自己的叶子工具。新草稿禁用且未发布，登记前后核对原有Agent的模型、工具、Skill、子Agent和提示词签名；不读取凭据给模型。"),
        ("table", (["组件", "实际职责"], [
          ["govmatter_open", "固定公开场景 开任务并固定版本"],
          ["govmatter_retrieve", "检索子Agent取得来源原文"],
          ["govmatter_verify", "核验子Agent检查引用覆盖及类型推导"],
          ["govmatter_state", "主控读取真实持久阶段并确认核验结果"],
          ["govmatter_finalize", "重查当前控制 返回原样候选结果"]])),
        ("h", "返回值数据流"),
        ("p", "v4解析MCP返回的JSON文本，从开任务结果取得真实ID，再构造两个子Agent任务JSON。单份完整代码动作保留开任务、状态及最终候选标记断言，最终字符串原样交给final_answer。宿主只允许json解析导入，不增加文件、额外网络或发布工具。"),
        ("p", f"本轮v4实际结果为{status}，耗时{native['duration_ms']/1000:.1f}秒，原生请求{native['native_run_requests']}次；底层model-attempt开始{native.get('model_attempt_begins',0)}次、rollback {native.get('model_attempt_rollbacks',0)}次。具体工具和结果以原始事件及报告为准。"),
        ("p", "v4活动固定1次与1200秒上限，主控至多2步和512输出token，子Agent各2步和256输出token。此前原2次、v2 1次、v3 1次的失败与计数全部保留。每次原生请求可能包含底层模型重试，不能将二者混算。"),
        ("p", "当前核验为确定性原文与类型程序检查，没有调用Jev或编造支持分数。本轮未调用云模型，也没有向政务系统提交申请。")]),
      ("更新发布与双端权限", [
        ("p", "更新链路区分资料层和正式业务知识：无变化只记录观察，来源新增或正文元数据变化建立候选，工程发布使用预期头版本进行竞争检查；正式业务知识需要实际维护者复核。用户已明确保留候选，因此本轮只发布工程预览。"),
        ("p", "群众登录后声明条件并查询，后端固定知识版本。一次性交接码用于将原条件、版本和引用交给工作人员；未分配工作人员不能读取他人案例。工作人员演示确认仍标记非正式业务批准，不把自动账号操作当真人核定。"),
        ("table", (["验收记录", "结果与适用范围"], [
          ["现有Nexent身份更新链路", "19项真实HTTP检查 原件更新工程及身份交接"],
          ["ZIP解压项目 新Python环境", "19项真实本地HTTP检查 随机工程账户"],
          ["OCR契约", "8项明确夹具检查 另有一次真实接口返回"],
          ["v4执行器预检", "4种工具夹具路径 无模型与真实子Agent调用"]])),
        ("p", "两组19项覆盖不同运行环境，但复用了开发场景，不能合计为38个独立业务案例。当前干净Python复现没有重建Nexent容器或在新机器重新加载模型，完整平台迁移仍需单独核查。"),
        ("h", "已知限制与独立评测"),
        ("p", "独立评测协议及空提交模板已准备，等待真实独立评测人、来源依据和未用于开发的问题组。先冻结参考，再运行预测，按原始问题组隔离同义改写；同时报告正确性、引用、条件、恰当停止和必要基线，不报告尚未取得的分数。"),
        ("p", "人工规则核定、OCR逐字和文类批准、独立业务评测、最终赛事模板及签收尚未完成。官方模板与当前赛事要求须另行核对，本说明不代替最终材料验收。")]),
      ("复现与证据索引", [
        ("p", "公开项目v1快照包含1448文件，ZIP为20,391,169字节；GitHub下载返回200，SHA256与本地一致。新建Python环境安装39个分发包后，从ZIP解压目录导入关键项目模块并启动双端HTTP，19项检查通过。包内快照固定原提交，v4等后续改动在当前分支另行提交，不能称旧ZIP已包含新修复。"),
        ("code", "python -m venv .venv\n.\\.venv\\Scripts\\python.exe -m pip install -r requirements.lock\n.\\.venv\\Scripts\\python.exe -m pip install -r requirements-knowledge.lock\n.\\.venv\\Scripts\\python.exe scripts/verify_matter_update_inputs.py\n.\\.venv\\Scripts\\python.exe scripts/reproduce_matter_project.py\n  --runtime .runtime/new-reproduction"),
        ("p", "上例最后两行表示同一条命令，执行时合并。复现会新建私有随机工程账户、保留隔离暂停终态并关闭自己的HTTP服务。超过归档新鲜度时应停止，不能改抓取日伪装最新；实际刷新和差异复核另走更新链路。"),
        ("table", (["证据", "位置"], [
          ["来源与规则", "docs/南京真实事项来源与规则复核包.md"],
          ["更新工程", "evidence/real-matter-closure/update-cycle/"],
          ["v4原生实录", "evidence/real-matter-closure/native-dataflow-v4/"],
          ["OCR原始返回", "data/real-matter-closure/nanjing/ocr-handoff/result-v1/"],
          ["解压与下载", "evidence/real-matter-closure/project-delivery/"],
          ["实际网页截图", "evidence/real-matter-closure/web/"],
          ["独立评测协议", "docs/南京事项独立评测协议.md"]])),
        ("p", "已归档网页截图属于较早工程版本及工程账户演示，不冒充v4模型或当前四来源更新的新截图。Windows五分钟恢复任务继续停用，只有用户明确开启才启用。"),
        ("p", "本说明事实基准提交为 " + commit + "。本地账户、运行库、令牌与权重不随公开交付；队友按部署说明在自己的环境登记模型与角色。")])
    ]
    doc = Document(); section = doc.sections[0]
    section.page_width = Inches(8.5); section.page_height = Inches(11)
    section.top_margin = Inches(.7); section.bottom_margin = Inches(.7)
    section.left_margin = Inches(.75); section.right_margin = Inches(.75)
    for name, size in (("Normal",11),("Title",21),("Heading 1",15),("Heading 2",12)):
        style=doc.styles[name]; style.font.name="宋体"; style.font.size=Pt(size); style.font.color.rgb=RGBColor(0,0,0)
        style.element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"),"宋体" if name=="Normal" else "黑体")
        style.paragraph_format.space_after=Pt(7); style.paragraph_format.line_spacing=1.16
    doc.styles["Normal"].paragraph_format.widow_control=True
    footer=section.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.CENTER
    footer.add_run("南京事项工程候选技术说明  ").font.size=Pt(9)
    field=OxmlElement("w:fldSimple");field.set(qn("w:instr"),"PAGE");footer._p.append(field)
    markdown=[]
    for index,(title,blocks) in enumerate(pages):
        if index:doc.add_page_break()
        doc.add_paragraph(title,"Title" if index==0 else "Heading 1");markdown.append("# "+title)
        for kind,value in blocks:
            if kind=="p":doc.add_paragraph(value);markdown.append(value)
            elif kind=="h":doc.add_paragraph(value,"Heading 2");markdown.append("## "+value)
            elif kind=="code":
                p=doc.add_paragraph();p.paragraph_format.line_spacing=1.05
                run=p.add_run(value);run.font.name="Consolas";run.font.size=Pt(9.5);markdown.append("```text\n"+value+"\n```")
            else:
                headers,rows=value;table=doc.add_table(rows=1,cols=len(headers));table.autofit=False
                widths=([2.25,.7,4.05] if headers==["来源","正文量","用途"] else [1.6,3.8,1.6] if len(headers)==3 else [2.05,4.95])
                for i,width in enumerate(widths):table.columns[i].width=Inches(width)
                for c,text in zip(table.rows[0].cells,headers):c.text=text
                repeat=OxmlElement("w:tblHeader");table.rows[0]._tr.get_or_add_trPr().append(repeat)
                for row in rows:
                    for c,text in zip(table.add_row().cells,row):c.text=text
                borders=OxmlElement("w:tblBorders")
                for side in ("top","left","bottom","right","insideH","insideV"):
                    e=OxmlElement("w:"+side);e.set(qn("w:val"),"single");e.set(qn("w:sz"),"4");e.set(qn("w:color"),"D9D9D9");borders.append(e)
                table._tbl.tblPr.append(borders)
                for ri,row in enumerate(table.rows):
                    for ci,cell in enumerate(row.cells):
                        cell.width=Inches(widths[ci]);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
                        shade=OxmlElement("w:shd");shade.set(qn("w:fill"),"D9E5EF" if ri==0 else "FFFFFF" if ri%2 else "F5F7FA");cell._tc.get_or_add_tcPr().append(shade)
                        for p in cell.paragraphs:
                            if (headers[ci]=="正文量" or headers[ci]=="状态"): p.alignment=WD_ALIGN_PARAGRAPH.CENTER
                            p.paragraph_format.space_before=Pt(4);p.paragraph_format.space_after=Pt(4)
                            p.paragraph_format.line_spacing=1.1
                            for run in p.runs:run.font.size=Pt(10);run.bold=ri==0
                doc.add_paragraph().paragraph_format.space_after=Pt(1)
                markdown.append("\n".join(["| "+" | ".join(headers)+" |", "| "+" | ".join(["---"]*len(headers))+" |",
                    *("| "+" | ".join(row)+" |" for row in rows)]))
        markdown.append("")
    for element in (doc.element, doc.styles.element):
        for border in list(element.xpath(".//w:pBdr")): border.getparent().remove(border)
    target=ROOT/"docs/南京真实事项技术说明.docx";doc.save(target)
    (ROOT/"docs/南京真实事项技术说明.md").write_text("\n\n".join(markdown).rstrip()+"\n",encoding="utf-8")
    print(json.dumps({"docx":str(target),"sha256":hashlib.sha256(target.read_bytes()).hexdigest(),"native_result":status,"human_review":"pending"},ensure_ascii=False))


if __name__=="__main__":main()
